import hashlib
from io import BytesIO
from pathlib import Path
from typing import Iterator
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import downloads
from app.database import SessionLocal
from app.main import app
from app.models import Tag, TagType, Work, WorkFile

client = TestClient(app)


@pytest.fixture
def upload_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> Iterator[Path]:
    monkeypatch.setattr(downloads, "UPLOAD_STORAGE_DIR", tmp_path)
    yield tmp_path


def create_work_with_file(
    content: bytes,
    storage_filename: str,
    original_filename: str = "My Original Book.epub",
) -> int:
    with SessionLocal() as db:
        work = Work(title="Download Test Work")
        work.files.append(
            WorkFile(
                original_filename=original_filename,
                storage_filename=storage_filename,
                file_format="EPUB",
                file_size=len(content),
                checksum=hashlib.sha256(content).hexdigest(),
            )
        )
        db.add(work)
        db.commit()
        db.refresh(work)
        return work.id


def delete_work(work_id: int) -> None:
    with SessionLocal() as db:
        work = db.get(Work, work_id)
        if work is not None:
            db.delete(work)
            db.commit()


def test_download_returns_original_bytes_and_filename(
    upload_directory: Path,
) -> None:
    content = b"original EPUB file bytes"
    storage_filename = "download-test-original.epub"
    (upload_directory / storage_filename).write_bytes(content)
    work_id = create_work_with_file(content, storage_filename)

    try:
        response = client.get(f"/api/works/{work_id}/download")

        assert response.status_code == 200
        assert response.content == content
        assert response.headers["content-type"] == "application/epub+zip"
        assert "filename*=utf-8''My%20Original%20Book.epub" in response.headers[
            "content-disposition"
        ]
        assert hashlib.sha256(response.content).hexdigest() == hashlib.sha256(
            content
        ).hexdigest()
    finally:
        delete_work(work_id)


def test_download_names_pdf_conversion_as_epub(upload_directory: Path) -> None:
    content = b"converted EPUB file bytes"
    storage_filename = "download-converted-pdf.epub"
    (upload_directory / storage_filename).write_bytes(content)
    work_id = create_work_with_file(
        content,
        storage_filename,
        original_filename="My Original Book.pdf",
    )

    try:
        response = client.get(f"/api/works/{work_id}/download")

        assert response.status_code == 200
        assert response.content == content
        assert "filename*=utf-8''My%20Original%20Book.epub" in response.headers[
            "content-disposition"
        ]
    finally:
        delete_work(work_id)


def test_bulk_download_contains_epubs_and_excel_metadata(
    upload_directory: Path,
) -> None:
    first_content = b"first EPUB"
    second_content = b"second EPUB"
    (upload_directory / "bulk-first.epub").write_bytes(first_content)
    (upload_directory / "bulk-second.epub").write_bytes(second_content)
    first_id = create_work_with_file(
        first_content,
        "bulk-first.epub",
        original_filename="First Story.epub",
    )
    second_id = create_work_with_file(
        second_content,
        "bulk-second.epub",
        original_filename="Second Story.epub",
    )
    with SessionLocal() as db:
        first_work = db.get(Work, first_id)
        assert first_work is not None
        first_work.title = "Bulk Export Story"
        first_work.original_author = "Export Author"
        first_work.description = "=1+1"
        first_work.tags.append(
            Tag(name=f"Bulk Export Fandom {first_id}", tag_type=TagType.FANDOM)
        )
        db.commit()

    try:
        response = client.post(
            "/api/works/download-bundle",
            json={"work_ids": [first_id, second_id]},
        )

        assert response.status_code == 200
        assert response.headers["content-type"] == "application/zip"
        with ZipFile(BytesIO(response.content)) as bundle:
            assert set(bundle.namelist()) == {
                "metadata.xlsx",
                f"fanfics/{first_id}-First_Story.epub",
                f"fanfics/{second_id}-Second_Story.epub",
            }
            assert bundle.read(f"fanfics/{first_id}-First_Story.epub") == first_content
            assert (
                bundle.read(f"fanfics/{second_id}-Second_Story.epub")
                == second_content
            )
            workbook = load_workbook(BytesIO(bundle.read("metadata.xlsx")))
            worksheet = workbook["Fic metadata"]
            headers = [cell.value for cell in worksheet[1]]
            first_row = dict(
                zip(headers, next(worksheet.iter_rows(min_row=2, values_only=True)))
            )
            assert first_row["Title"] == "Bulk Export Story"
            assert first_row["Author"] == "Export Author"
            assert first_row["Description"] == "'=1+1"
            assert first_row["Fandoms"] == f"Bulk Export Fandom {first_id}"
            assert first_row["Original filename"] == "First Story.epub"
            assert worksheet.max_row == 3
    finally:
        delete_work(first_id)
        delete_work(second_id)


def test_bulk_download_rejects_missing_or_duplicate_ids() -> None:
    missing_response = client.post(
        "/api/works/download-bundle",
        json={"work_ids": [2147483647]},
    )
    duplicate_response = client.post(
        "/api/works/download-bundle",
        json={"work_ids": [1, 1]},
    )

    assert missing_response.status_code == 404
    assert duplicate_response.status_code == 422


def test_download_returns_404_for_missing_work() -> None:
    response = client.get("/api/works/2147483647/download")

    assert response.status_code == 404
    assert response.json()["detail"] == "Work not found"


def test_download_returns_404_when_stored_file_is_missing(
    upload_directory: Path,
) -> None:
    work_id = create_work_with_file(b"missing file", "missing.epub")

    try:
        response = client.get(f"/api/works/{work_id}/download")

        assert response.status_code == 404
        assert response.json()["detail"] == "Stored file not found"
    finally:
        delete_work(work_id)


def test_download_does_not_allow_stored_path_to_escape_storage(
    upload_directory: Path, tmp_path: Path
) -> None:
    outside_file = tmp_path.parent / "outside-download-test.epub"
    outside_file.write_bytes(b"must not be downloadable")
    work_id = create_work_with_file(
        outside_file.read_bytes(),
        f"../{outside_file.name}",
        original_filename="outside.epub",
    )

    try:
        response = client.get(f"/api/works/{work_id}/download")

        assert response.status_code == 404
        assert response.json()["detail"] == "Stored file not found"
    finally:
        delete_work(work_id)
        outside_file.unlink(missing_ok=True)
