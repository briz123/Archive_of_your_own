from datetime import datetime, timezone
from enum import Enum
from io import BytesIO
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import TagType, Work

WORK_COLUMNS = [
    ("Work ID", 12),
    ("Fandom group", 28),
    ("Group size", 12),
    ("Work title", 36),
    ("Author", 24),
    ("Summary", 60),
    ("Fandoms", 36),
    ("Characters", 36),
    ("Relationships", 36),
    ("Additional tags", 40),
    ("Rating", 24),
    ("Archive warnings", 40),
    ("Category", 18),
    ("Language", 16),
    ("Publisher", 24),
    ("Identifier", 28),
    ("Subject", 36),
    ("Creator", 24),
    ("Producer", 24),
    ("Publication date", 18),
    ("Creation date", 18),
    ("Modification date", 18),
    ("Word count", 14),
    ("Chapter count", 14),
    ("Completion status", 20),
    ("Published at", 20),
    ("Date added", 20),
    ("Date updated", 20),
    ("File formats", 18),
    ("Original filenames", 40),
    ("File sizes (bytes)", 20),
]
TAG_COLUMNS = [
    ("Work ID", 12),
    ("Work title", 40),
    ("Tag name", 44),
    ("Tag type", 20),
]
TAG_TYPE_ORDER = {
    TagType.FANDOM: 0,
    TagType.CHARACTER: 1,
    TagType.RELATIONSHIP: 2,
    TagType.ADDITIONAL: 3,
}
BLACK_CHERRY = "FF4B1022"
WINE = "FF78233A"
WINE_RED = "FF96384C"
IVORY = "FFFAF7F5"
BLUSH = "FFF8F0EC"
PALE_ROSE = "FFF4E9E6"
ROSE = "FFEADCE3"
SOFT_ROSE = "FFF0DDE1"
WHITE = "FFFFFFFF"
FANDOM_ACCENTS = (PALE_ROSE, BLUSH, ROSE, SOFT_ROSE)
TAG_TYPE_FILLS = {
    TagType.FANDOM: PALE_ROSE,
    TagType.CHARACTER: ROSE,
    TagType.RELATIONSHIP: SOFT_ROSE,
    TagType.ADDITIONAL: BLUSH,
}
DATE_FORMAT = "yyyy-mm-dd hh:mm"


def _value(value):
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def _work_tags(work: Work):
    return [tag for tag in work.tags if tag.user_id == work.user_id]


def _fandoms(work: Work):
    return sorted(
        (tag for tag in _work_tags(work) if tag.tag_type == TagType.FANDOM),
        key=lambda tag: tag.name.casefold(),
    )


def _group_works(works: Iterable[Work]):
    groups = {}
    for work in works:
        fandoms = _fandoms(work)
        if len(fandoms) > 1:
            group_type, group_name = "crossover", "Crossover"
        elif len(fandoms) == 1:
            group_type, group_name = "fandom", fandoms[0].name
        else:
            group_type, group_name = "uncategorized", "Uncategorized"
        key = (group_type, group_name)
        groups.setdefault(key, []).append(work)

    group_order = {"fandom": 0, "crossover": 1, "uncategorized": 2}
    return sorted(
        (
            (
                group_type,
                group_name,
                sorted(
                    group_works,
                    key=lambda work: (work.title.casefold(), work.id),
                ),
            )
            for (group_type, group_name), group_works in groups.items()
        ),
        key=lambda group: (
            group_order[group[0]],
            group[1].casefold(),
            group[1],
        ),
    )


def _join_tag_names(work: Work, tag_type: TagType) -> str:
    return "; ".join(
        tag.name
        for tag in sorted(
            (tag for tag in _work_tags(work) if tag.tag_type == tag_type),
            key=lambda tag: tag.name.casefold(),
        )
    )


def _style_header(worksheet: Worksheet) -> None:
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=BLACK_CHERRY)
        cell.font = Font(name="Calibri", bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(
            bottom=Side(style="medium", color=ROSE),
        )
    worksheet.row_dimensions[1].height = 30
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False


def _style_data_row(worksheet: Worksheet, row_number: int, banded: bool) -> None:
    fill = PatternFill("solid", fgColor=IVORY) if banded else PatternFill()
    for cell in worksheet[row_number]:
        cell.fill = fill
        cell.font = Font(name="Calibri", size=10, color=BLACK_CHERRY)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    title_cell = worksheet.cell(row_number, 4)
    title_cell.font = Font(
        name="Georgia",
        size=11,
        bold=True,
        color=BLACK_CHERRY,
    )


def _style_group_heading(
    worksheet: Worksheet,
    row_number: int,
    group_type: str,
    fandom_index: int,
) -> None:
    if group_type == "crossover":
        accent = BLACK_CHERRY
        text_color = WHITE
    elif group_type == "uncategorized":
        accent = ROSE
        text_color = BLACK_CHERRY
    else:
        accent = FANDOM_ACCENTS[fandom_index % len(FANDOM_ACCENTS)]
        text_color = BLACK_CHERRY

    for cell in worksheet[row_number][1:3]:
        cell.fill = PatternFill("solid", fgColor=accent)
        cell.font = Font(name="Calibri", bold=True, color=text_color)
        cell.alignment = Alignment(vertical="center")
    worksheet.cell(row_number, 2).border = Border(
        left=Side(
            style="medium",
            color=BLACK_CHERRY if group_type == "crossover" else accent,
        ),
    )


def _configure_data_sheet(
    worksheet: Worksheet,
    columns: list[tuple[str, int]],
) -> None:
    worksheet.append([name for name, _ in columns])
    for index, (_, width) in enumerate(columns, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    _style_header(worksheet)


def _work_values(work: Work, group_name: str, group_size: Optional[int]) -> list:
    files = work.files
    return [
        work.id,
        group_name,
        group_size,
        work.title,
        work.original_author,
        work.description,
        _join_tag_names(work, TagType.FANDOM),
        _join_tag_names(work, TagType.CHARACTER),
        _join_tag_names(work, TagType.RELATIONSHIP),
        _join_tag_names(work, TagType.ADDITIONAL),
        _value(work.rating),
        "; ".join(sorted((_value(warning.name) for warning in work.warnings), key=str.casefold)),
        _value(work.category),
        work.language,
        work.publisher,
        work.identifier,
        work.subject,
        work.creator,
        work.producer,
        work.publication_date,
        work.creation_date,
        work.modification_date,
        work.word_count,
        work.chapter_count,
        work.completion_status,
        _value(work.published_at),
        _value(work.created_at),
        _value(work.updated_at),
        "; ".join(sorted({file.file_format for file in files}, key=str.casefold)),
        "; ".join(file.original_filename for file in files),
        "; ".join(str(file.file_size) for file in files),
    ]


def _add_all_works_sheet(workbook: Workbook, works: list[Work]) -> None:
    worksheet = workbook.create_sheet("All Works")
    worksheet.sheet_properties.tabColor = WINE
    _configure_data_sheet(worksheet, WORK_COLUMNS)

    fandom_index = 0
    for group_type, group_name, group_works in _group_works(works):
        for index, work in enumerate(group_works):
            worksheet.append(
                _work_values(
                    work,
                    group_name,
                    len(group_works) if index == 0 else None,
                )
            )
            row_number = worksheet.max_row
            _style_data_row(worksheet, row_number, index % 2 == 1)
            if index == 0:
                _style_group_heading(
                    worksheet,
                    row_number,
                    group_type,
                    fandom_index,
                )
            if group_type == "fandom" and index == 0:
                fandom_index += 1
            for column_name in ("Published at", "Date added", "Date updated"):
                column_number = next(
                    index
                    for index, (name, _) in enumerate(WORK_COLUMNS, start=1)
                    if name == column_name
                )
                worksheet.cell(row_number, column_number).number_format = DATE_FORMAT

    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(WORK_COLUMNS))}{worksheet.max_row}"
    )


def _add_tags_sheet(workbook: Workbook, works: list[Work]) -> None:
    worksheet = workbook.create_sheet("Tags")
    worksheet.sheet_properties.tabColor = WINE_RED
    _configure_data_sheet(worksheet, TAG_COLUMNS)

    tag_rows = [
        (
            work.id,
            work.title,
            tag.name,
            tag.tag_type.value,
            TAG_TYPE_ORDER[tag.tag_type],
        )
        for work in works
        for tag in _work_tags(work)
    ]
    tag_rows.sort(
        key=lambda row: (
            row[1].casefold(),
            row[4],
            row[2].casefold(),
            row[0],
        )
    )
    for index, tag_row in enumerate(tag_rows):
        worksheet.append(tag_row[:4])
        row_number = worksheet.max_row
        _style_data_row(worksheet, row_number, index % 2 == 1)
        worksheet.cell(row_number, 2).font = Font(
            name="Georgia",
            size=10,
            bold=True,
            color=BLACK_CHERRY,
        )
        type_cell = worksheet.cell(row_number, 4)
        type_cell.fill = PatternFill(
            "solid",
            fgColor=TAG_TYPE_FILLS[TagType(tag_row[3])],
        )
        type_cell.font = Font(name="Calibri", size=10, bold=True, color=BLACK_CHERRY)
    worksheet.auto_filter.ref = f"A1:D{worksheet.max_row}"


def _add_read_me_sheet(
    workbook: Workbook,
    works: list[Work],
    generated_at: datetime,
) -> None:
    worksheet = workbook.active
    worksheet.title = "Read Me"
    worksheet.sheet_properties.tabColor = BLACK_CHERRY
    worksheet.sheet_view.showGridLines = False
    for column in "BCDEFG":
        worksheet.column_dimensions[column].width = 16
    worksheet.column_dimensions["A"].width = 3
    worksheet.column_dimensions["H"].width = 3

    worksheet.merge_cells("B2:G2")
    worksheet["B2"] = "ARCHIVE OF YOUR OWN  /  AYO"
    worksheet["B2"].font = Font(name="Calibri", size=10, bold=True, color=WINE)
    worksheet["B2"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[2].height = 24

    worksheet.merge_cells("B3:G3")
    worksheet["B3"] = "My Personal Library"
    worksheet["B3"].font = Font(name="Georgia", size=26, bold=True, color=BLACK_CHERRY)
    worksheet["B3"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[3].height = 42

    worksheet.merge_cells("B4:G4")
    worksheet["B4"] = (
        "A considered catalog of the stories in your personal collection."
    )
    worksheet["B4"].font = Font(name="Calibri", size=11, color=BLACK_CHERRY)
    worksheet["B4"].alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[4].height = 25

    export_date = _value(generated_at)
    worksheet.merge_cells("B5:G5")
    worksheet["B5"] = (
        f"EXPORTED  ·  {export_date.strftime('%B')} "
        f"{export_date.day}, {export_date.year}"
    )
    worksheet["B5"].font = Font(name="Calibri", size=9, bold=True, color=WINE)
    worksheet["B5"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[5].height = 22

    fandom_count = len({
        tag.name.casefold()
        for work in works
        for tag in _fandoms(work)
    })
    crossover_count = sum(
        len(group_works)
        for group_type, _, group_works in _group_works(works)
        if group_type == "crossover"
    )
    stats = (
        ("B", "C", "TOTAL WORKS", len(works), "cataloged in this export", PALE_ROSE),
        ("D", "E", "FANDOMS", fandom_count, "distinct fandoms", BLUSH),
        ("F", "G", "CROSSOVER WORKS", crossover_count, "two or more fandoms", ROSE),
    )
    for first_column, last_column, label, value, detail, accent in stats:
        for row_number in range(7, 10):
            worksheet.merge_cells(
                start_row=row_number,
                start_column=ord(first_column) - ord("A") + 1,
                end_row=row_number,
                end_column=ord(last_column) - ord("A") + 1,
            )
            for column_number in range(
                ord(first_column) - ord("A") + 1,
                ord(last_column) - ord("A") + 2,
            ):
                worksheet.cell(row_number, column_number).fill = PatternFill(
                    "solid",
                    fgColor=accent,
                )
        label_cell = worksheet[f"{first_column}7"]
        label_cell.value = label
        label_cell.font = Font(name="Calibri", size=9, bold=True, color=WINE)
        label_cell.alignment = Alignment(vertical="center")
        value_cell = worksheet[f"{first_column}8"]
        value_cell.value = value
        value_cell.font = Font(name="Georgia", size=22, bold=True, color=BLACK_CHERRY)
        value_cell.alignment = Alignment(vertical="center")
        detail_cell = worksheet[f"{first_column}9"]
        detail_cell.value = detail
        detail_cell.font = Font(name="Calibri", size=9, color=BLACK_CHERRY)
        detail_cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[7].height = 20
    worksheet.row_dimensions[8].height = 38
    worksheet.row_dimensions[9].height = 25

    worksheet.merge_cells("B11:G11")
    worksheet["B11"] = "A guide to your workbook"
    worksheet["B11"].font = Font(name="Georgia", size=16, bold=True, color=BLACK_CHERRY)
    worksheet["B11"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[11].height = 32
    guide_rows = (
        (
            "B12:C12",
            "D12:G12",
            "ALL WORKS",
            "One row per work, arranged in fandom sections. Crossovers appear once in their own section; works without a fandom tag are listed as Uncategorized.",
            PALE_ROSE,
        ),
        (
            "B13:C13",
            "D13:G13",
            "TAGS",
            "A filterable record of every tag, with its work ID, title, name, and tag type.",
            BLUSH,
        ),
        (
            "B14:C14",
            "D14:G14",
            "READ ME",
            "This overview includes collection totals and explains how the workbook is organized.",
            SOFT_ROSE,
        ),
    )
    for label_range, description_range, label, description, accent in guide_rows:
        worksheet.merge_cells(label_range)
        worksheet.merge_cells(description_range)
        label_cell = worksheet[label_range.split(":")[0]]
        label_cell.value = label
        label_cell.fill = PatternFill("solid", fgColor=accent)
        label_cell.font = Font(name="Calibri", size=9, bold=True, color=WINE)
        label_cell.alignment = Alignment(vertical="center")
        description_cell = worksheet[description_range.split(":")[0]]
        description_cell.value = description
        description_cell.font = Font(name="Calibri", size=10, color=BLACK_CHERRY)
        description_cell.alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.row_dimensions[description_cell.row].height = 38

    worksheet.merge_cells("B16:G16")
    worksheet["B16"] = (
        "Summaries and tag lists are wrapped for reading; use the table filters "
        "on All Works and Tags to browse or sort your catalog."
    )
    worksheet["B16"].font = Font(name="Calibri", size=10, color=BLACK_CHERRY, italic=True)
    worksheet["B16"].alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[16].height = 32

    work_definitions = [
        ("Work ID", "AYO's ID for the work."),
        ("Fandom group", "Section grouping based on the number of fandom tag relationships."),
        ("Group size", "Number of works in this group, shown on its first work row."),
        ("Work title", "The saved work title."),
        ("Author", "Original author, when available."),
        ("Summary", "Saved work description."),
        ("Fandoms", "All individual fandom tag names for the work."),
        ("Characters", "All individual character tag names for the work."),
        ("Relationships", "All individual relationship tag names for the work."),
        ("Additional tags", "All individual additional tag names for the work."),
        ("Rating", "Saved rating."),
        ("Archive warnings", "All saved archive warning names."),
        ("Category", "Saved work category."),
        ("Language", "Saved language."),
        ("Publisher", "Saved publisher."),
        ("Identifier", "Saved identifier."),
        ("Subject", "Saved subject."),
        ("Creator", "Saved creator."),
        ("Producer", "Saved producer."),
        ("Publication date", "Original publication date metadata."),
        ("Creation date", "Original creation date metadata."),
        ("Modification date", "Original modification date metadata."),
        ("Word count", "Saved word count, when available."),
        ("Chapter count", "Saved chapter count, when available."),
        ("Completion status", "Saved completion status."),
        ("Published at", "Saved publication timestamp."),
        ("Date added", "When the work was added to AYO."),
        ("Date updated", "When the work metadata was last updated."),
        ("File formats", "Formats of associated local-file records."),
        ("Original filenames", "Filenames from associated local-file records."),
        ("File sizes (bytes)", "Sizes recorded for associated local-file records."),
    ]
    tag_definitions = [
        ("Work ID", "AYO's ID for the work associated with this tag."),
        ("Work title", "Saved work title."),
        ("Tag name", "One individual saved tag name."),
        ("Tag type", "Tag category: fandom, character, relationship, or additional."),
    ]
    for section_row, title, definitions, accent in (
        (18, "All Works column guide", work_definitions, PALE_ROSE),
        (52, "Tags column guide", tag_definitions, BLUSH),
    ):
        worksheet.merge_cells(
            start_row=section_row,
            start_column=2,
            end_row=section_row,
            end_column=7,
        )
        section_cell = worksheet.cell(section_row, 2)
        section_cell.value = title
        section_cell.font = Font(name="Georgia", size=15, bold=True, color=BLACK_CHERRY)
        section_cell.alignment = Alignment(vertical="center")
        worksheet.row_dimensions[section_row].height = 30

        header_row = section_row + 1
        for cell_range in (f"B{header_row}:C{header_row}", f"D{header_row}:G{header_row}"):
            worksheet.merge_cells(cell_range)
        for cell in (worksheet.cell(header_row, 2), worksheet.cell(header_row, 4)):
            cell.fill = PatternFill("solid", fgColor=accent)
            cell.font = Font(name="Calibri", size=9, bold=True, color=WINE)
            cell.alignment = Alignment(vertical="center")
        worksheet.cell(header_row, 2).value = "COLUMN"
        worksheet.cell(header_row, 4).value = "MEANING"
        worksheet.row_dimensions[header_row].height = 22

        for index, (column_name, meaning) in enumerate(definitions):
            row_number = header_row + index + 1
            worksheet.merge_cells(
                start_row=row_number,
                start_column=2,
                end_row=row_number,
                end_column=3,
            )
            worksheet.merge_cells(
                start_row=row_number,
                start_column=4,
                end_row=row_number,
                end_column=7,
            )
            row_fill = IVORY if index % 2 else WHITE
            for column_number in range(2, 8):
                worksheet.cell(row_number, column_number).fill = PatternFill(
                    "solid",
                    fgColor=row_fill,
                )
            name_cell = worksheet.cell(row_number, 2)
            name_cell.value = column_name
            name_cell.font = Font(name="Calibri", size=9, bold=True, color=BLACK_CHERRY)
            name_cell.alignment = Alignment(vertical="top", wrap_text=True)
            meaning_cell = worksheet.cell(row_number, 4)
            meaning_cell.value = meaning
            meaning_cell.font = Font(name="Calibri", size=9, color=BLACK_CHERRY)
            meaning_cell.alignment = Alignment(vertical="top", wrap_text=True)
            worksheet.row_dimensions[row_number].height = 25


def build_library_workbook(
    works: list[Work],
    generated_at: Optional[datetime] = None,
) -> bytes:
    generated_at = generated_at or datetime.now(timezone.utc)
    workbook = Workbook()
    workbook.properties.creator = "Archive of Your Own"
    workbook.properties.created = _value(generated_at)
    _add_read_me_sheet(workbook, works, generated_at)
    _add_all_works_sheet(workbook, works)
    _add_tags_sheet(workbook, works)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
